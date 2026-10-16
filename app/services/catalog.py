"""Catalogue : import/export CSV-XLSX validé, collections par règles, recherche et filtres, données de démo."""
import base64
import csv
import io
import re
import unicodedata

from sqlalchemy import select

from ..models import Product, SizeGuide, Supplier, Variant
from .audit import audit
from .jobs import handler

ALIASES = {
    "handle": ["handle", "identifiant"], "title": ["title", "titre", "nom", "name"], "brand": ["brand", "marque"],
    "model": ["model", "modele", "modèle"], "category": ["category", "categorie", "catégorie"], "usage": ["usage", "activite"],
    "gender": ["gender", "genre"], "color": ["color", "couleur", "colour"], "size": ["size", "pointure", "taille"],
    "sku": ["sku", "reference", "référence", "ref"], "price": ["price", "prix", "prix_ttc"],
    "cost": ["cost", "cout", "coût", "prix_achat", "cost_price"], "stock": ["stock", "quantite", "quantité", "qty"],
    "description": ["description"], "image_url": ["image_url", "image", "photo"], "image_license": ["image_license", "licence_image", "droits_image"],
    "supplier": ["supplier", "fournisseur"], "origin": ["origin", "provenance"], "authenticity": ["authenticity", "authenticite", "authenticité"],
    "tags": ["tags", "etiquettes"],
}
ATTR_PREFIX = "attr_"
EXPORT_COLS = ["handle", "title", "brand", "model", "category", "usage", "gender", "color", "size", "sku", "price",
               "cost", "stock", "description", "image_url", "image_license", "supplier", "origin", "authenticity", "tags"]
# + une colonne « attr_<clé> » par caractéristique (attr_movement, attr_metal, attr_ingredients…)


def slugify(s: str) -> str:
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:150] or "produit"


def _money(v) -> int | None:
    if v is None or str(v).strip() == "":
        return None
    s = str(v).replace("€", "").replace(" ", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    s = s.replace(",", ".")
    val = float(s)
    if val < 0:
        raise ValueError("montant négatif")
    return int(round(val * 100))


def read_rows(filename: str, content: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [{header[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(header)} for r in rows[1:]
                if any(v not in (None, "") for v in r)]
    text = content.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    return [r for r in csv.DictReader(io.StringIO(text), dialect=dialect) if any((v or "").strip() for v in r.values())]


def normalise_row(raw: dict) -> dict:
    low = {slugify(str(k)).replace("-", "_"): v for k, v in raw.items()}
    out = {"attributes": {k[len(ATTR_PREFIX):]: str(v).strip() for k, v in low.items()
                          if k.startswith(ATTR_PREFIX) and str(v).strip() != ""}}
    for field, names in ALIASES.items():
        for n in names:
            key = slugify(n).replace("-", "_")
            if key in low and str(low[key]).strip() != "":
                out[field] = str(low[key]).strip()
                break
    return out


def normalise_category(value: str) -> str:
    from . import verticals
    v = slugify(value)
    if v in verticals.PROFILES:
        return v
    if v in verticals.ALIASES:
        return verticals.ALIASES[v]
    found = verticals.detect(value)
    return found[0] if found else v


def import_rows(db, shop, rows: list[dict], user_id=None, dry_run=False) -> dict:
    report = {"rows": len(rows), "products_created": 0, "variants_created": 0, "variants_updated": 0,
              "errors": [], "warnings": []}
    seen = set()
    suppliers = {s.name: s for s in db.scalars(select(Supplier).where(Supplier.shop_id == shop.id))}
    for i, raw in enumerate(rows, start=2):  # ligne 1 = en-têtes
        r = normalise_row(raw)
        def err(field, msg):
            report["errors"].append({"row": i, "field": field, "message": msg})
        sku = r.get("sku")
        if not sku:
            err("sku", "SKU obligatoire"); continue
        if sku in seen:
            err("sku", f"SKU en double dans le fichier : {sku}"); continue
        seen.add(sku)
        try:
            price = _money(r.get("price"))
            cost = _money(r.get("cost"))
        except ValueError:
            err("price", "Prix ou coût invalide"); continue
        try:
            stock = int(float(r.get("stock", 0) or 0))
            if stock < 0:
                raise ValueError
        except ValueError:
            err("stock", "Stock invalide (entier ≥ 0 attendu)"); continue
        if price is None:
            report["warnings"].append({"row": i, "field": "price", "message": "Sans prix : la variante ne sera pas vendable"})
        if cost is None:
            report["warnings"].append({"row": i, "field": "cost", "message": "Coût d'achat manquant : marge non calculable"})
        if price is not None and cost is not None and cost >= price:
            report["warnings"].append({"row": i, "field": "price", "message": "Coût ≥ prix TTC : vente à perte"})
        img = r.get("image_url", "")
        if img and not img.startswith("https://"):
            err("image_url", "URL d'image en https obligatoire"); continue
        if img and not r.get("image_license"):
            report["warnings"].append({"row": i, "field": "image_license",
                                       "message": "Droits d'utilisation de l'image non déclarés : l'image sera masquée tant qu'ils ne sont pas renseignés"})
        variant = db.scalar(select(Variant).where(Variant.shop_id == shop.id, Variant.sku == sku))
        if variant is None:
            handle = r.get("handle") or slugify(r.get("title", "") + " " + r.get("color", ""))
            product = db.scalar(select(Product).where(Product.shop_id == shop.id, Product.handle == handle))
            if product is None:
                if not r.get("title"):
                    err("title", "Titre obligatoire pour un nouveau produit"); continue
                product = Product(shop_id=shop.id, handle=handle, title=r["title"])
                if not dry_run:
                    db.add(product)
                report["products_created"] += 1
            variant = Variant(shop_id=shop.id, product=product, sku=sku)
            if not dry_run:
                db.add(variant)
            report["variants_created"] += 1
        else:
            product = variant.product
            report["variants_updated"] += 1
            if stock < variant.reserved:
                err("stock", f"Stock {stock} inférieur aux quantités réservées ({variant.reserved})"); continue
        if dry_run:
            continue
        for f in ("brand", "model", "description", "origin", "authenticity"):
            if r.get(f):
                setattr(product, f, r[f])
        if r.get("category"):
            product.category = normalise_category(r["category"])
        if r["attributes"]:
            product.attributes = {**(product.attributes or {}), **r["attributes"]}
        if r.get("usage"):
            product.usage = slugify(r["usage"])
        if r.get("gender"):
            product.gender = slugify(r["gender"])
        if r.get("tags"):
            product.tags = [t.strip() for t in r["tags"].split("|") if t.strip()]
        if img and not any(x["url"] == img for x in product.images or []):
            product.images = (product.images or []) + [{"url": img, "alt": product.title, "license": r.get("image_license", "")}]
        if r.get("supplier"):
            s = suppliers.get(r["supplier"])
            if s is None:
                s = suppliers[r["supplier"]] = Supplier(shop_id=shop.id, name=r["supplier"])
                db.add(s)
            product.supplier = s
        variant.size, variant.color = r.get("size", variant.size), r.get("color", variant.color)
        if price is not None:
            variant.price_cents = price
        if cost is not None:
            variant.cost_cents = cost
        variant.stock = stock
        db.flush()
    if not dry_run:
        audit(db, shop.id, user_id, "catalog.import", "catalog", "", after={k: v for k, v in report.items() if k != "warnings"})
    return report


@handler("catalog_import")
def _job_import(db, job):
    from ..models import Shop
    shop = db.get(Shop, job.shop_id)
    rows = read_rows(job.payload["filename"], base64.b64decode(job.payload["content_b64"]))
    return import_rows(db, shop, rows, job.payload.get("user_id"))


def export_csv(db, shop) -> str:
    buf = io.StringIO()
    prods = db.scalars(select(Product).where(Product.shop_id == shop.id).order_by(Product.id)).all()
    attr_cols = sorted({ATTR_PREFIX + k for p in prods for k in (p.attributes or {})})
    w = csv.DictWriter(buf, fieldnames=EXPORT_COLS + attr_cols)
    w.writeheader()
    for p in prods:
        img = (p.images or [{}])[0] if p.images else {}
        for v in p.variants:
            w.writerow({"handle": p.handle, "title": p.title, "brand": p.brand, "model": p.model, "category": p.category,
                        "usage": p.usage, "gender": p.gender, "color": v.color, "size": v.size, "sku": v.sku,
                        "price": f"{v.price_cents/100:.2f}" if v.price_cents is not None else "",
                        "cost": f"{v.cost_cents/100:.2f}" if v.cost_cents is not None else "", "stock": v.stock,
                        "description": p.description, "image_url": img.get("url", ""), "image_license": img.get("license", ""),
                        "supplier": p.supplier.name if p.supplier else "", "origin": p.origin, "authenticity": p.authenticity,
                        "tags": "|".join(p.tags or []), **{ATTR_PREFIX + k: v for k, v in (p.attributes or {}).items()}})
    return buf.getvalue()


# ---------------------------------------------------------------- vitrine : collections, recherche, filtres
def collection_match(product: Product, rule: dict) -> bool:
    f, v = rule.get("field"), rule.get("value")
    if rule.get("category") and product.category != rule["category"]:
        return False
    if f == "all":
        return True
    if f == "tag":
        return v in (product.tags or [])
    return slugify(str(field_value(product, f) or "")) == slugify(v or "")


def field_value(product: Product, field: str):
    if field.startswith("attr:"):
        return (product.attributes or {}).get(field[5:], "")
    return getattr(product, field, "")


def sellable(db, shop_id: int):
    return db.scalars(select(Product).where(Product.shop_id == shop_id, Product.status == "active").order_by(Product.id)).all()


def min_price(p: Product):
    prices = [v.price_cents for v in p.variants if v.price_cents is not None]
    return min(prices) if prices else None


def search(products, q="", size="", color="", usage="", gender="", brand="", price_min=None, price_max=None,
           in_stock=False, sort="relevance", attrs=None, category=""):
    q = slugify(q) if q else ""
    attrs = {k: v for k, v in (attrs or {}).items() if v}
    out = []
    for p in products:
        if category and p.category != category:
            continue
        if any(slugify(str((p.attributes or {}).get(k, ""))) != slugify(v) for k, v in attrs.items()):
            continue
        if q:
            hay = slugify(" ".join([p.title, p.brand, p.model, p.usage, p.gender, p.description, " ".join(p.tags or []),
                                    " ".join(v.color for v in p.variants), " ".join(str(x) for x in (p.attributes or {}).values())]))
            if not all(tok in hay for tok in q.split("-")):
                continue
        vs = [v for v in p.variants if v.price_cents is not None]
        if size:
            vs = [v for v in vs if v.size == size]
        if color:
            vs = [v for v in vs if slugify(v.color) == slugify(color)]
        if in_stock:
            vs = [v for v in vs if v.available > 0]
        if price_min is not None:
            vs = [v for v in vs if v.price_cents >= price_min]
        if price_max is not None:
            vs = [v for v in vs if v.price_cents <= price_max]
        if not vs:
            continue
        if usage and slugify(p.usage) != slugify(usage):
            continue
        if gender and slugify(p.gender) != slugify(gender):
            continue
        if brand and slugify(p.brand) != slugify(brand):
            continue
        out.append(p)
    if sort == "price_asc":
        out.sort(key=lambda p: min_price(p) or 0)
    elif sort == "price_desc":
        out.sort(key=lambda p: -(min_price(p) or 0))
    elif sort == "new":
        out.sort(key=lambda p: -p.id)
    return out


def facets(products) -> dict:
    sizes, colors, usages, genders, brands, cats = set(), set(), set(), set(), set(), set()
    attrs: dict[str, set] = {}
    for p in products:
        for v in p.variants:
            if v.size:
                sizes.add(v.size)
            if v.color:
                colors.add(v.color)
        for s_, val in ((usages, p.usage), (genders, p.gender), (brands, p.brand), (cats, p.category)):
            if val:
                s_.add(val)
        for k, val in (p.attributes or {}).items():
            if val not in (None, ""):
                attrs.setdefault(k, set()).add(str(val))

    def num(x):
        try:
            return (0, float(x.replace(",", ".")))
        except ValueError:
            order = ["XXS", "XS", "S", "M", "L", "XL", "XXL"]
            return (1, order.index(x) if x in order else 99, x)
    return {"size": sorted(sizes, key=num), "color": sorted(colors), "usage": sorted(usages),
            "gender": sorted(genders), "brand": sorted(brands), "category": sorted(cats),
            "attrs": {k: sorted(v, key=num) for k, v in attrs.items()}}


def size_guide_for(db, shop_id: int, product: Product):
    """Guide du modèle si disponible, sinon guide générique de la marque (même verticale)."""
    guides = db.scalars(select(SizeGuide).where(SizeGuide.shop_id == shop_id, SizeGuide.brand == product.brand)).all()
    guides = [g for g in guides if not g.vertical or not product.category or g.vertical == product.category]
    exact = [g for g in guides if g.model and product.model and slugify(g.model) == slugify(product.model)]
    if exact:
        return exact[0], "model"
    generic = [g for g in guides if not g.model]
    return (generic[0], "brand") if generic else (None, None)


def seed_demo(db, shop, vertical_keys=None):
    """Produits de DÉMONSTRATION : marqués is_demo, badge visible, exclus des statistiques réelles."""
    from . import verticals
    for vk in vertical_keys or ["footwear"]:
        vp = verticals.get(vk)
        demo = vp["demo"]
        for i, (title, usage, gender, color, price, cost, sizes, attrs) in enumerate(demo["products"]):
            handle = slugify(title)
            if db.scalar(select(Product).where(Product.shop_id == shop.id, Product.handle == handle)):
                continue
            p = Product(shop_id=shop.id, handle=handle, title=title, brand=demo["brand"], model=title.split(" démo ")[-1],
                        usage=usage, gender=gender, category=vp["key"], attributes=dict(attrs), is_demo=True,
                        description="Produit fictif de démonstration. Remplacez-le par votre catalogue réel.",
                        origin="Démonstration — aucune provenance réelle")
            db.add(p)
            for j, sz in enumerate(sizes):
                db.add(Variant(shop_id=shop.id, product=p, sku=f"DEMO-{vp['key'][:3].upper()}-{i}-{slugify(sz)}", size=sz, color=color,
                               price_cents=price * 100, cost_cents=cost * 100, stock=(j * 3 + i) % 6 if len(sizes) > 1 else 4))
        if demo.get("guide") and vp.get("fit"):
            db.add(SizeGuide(shop_id=shop.id, brand=demo["brand"], model="", vertical=vp["key"], measure_label=vp["fit"]["measure_label"],
                             source_url="", notes="Guide FICTIF de démonstration.", rows=demo["guide"]))
    db.flush()


def purge_demo(db, shop):
    for p in db.scalars(select(Product).where(Product.shop_id == shop.id, Product.is_demo.is_(True))):
        db.delete(p)
    for g in db.scalars(select(SizeGuide).where(SizeGuide.shop_id == shop.id, SizeGuide.notes == "Guide FICTIF de démonstration.")):
        db.delete(g)
    db.flush()
